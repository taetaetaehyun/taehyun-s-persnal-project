# -*- coding: utf-8 -*-
"""
SPC Chart 엑셀 템플릿 생성기 (엑셀 2016 호환, 매크로 없음)

공용 폴더의 두 CSV 를 파워 쿼리로 '결과기록' / '측정데이터' 시트에 불러오면
'SPC Chart' 시트가 선택한 기록(기록ID)의 측정 Data 표, 통계치 Summary,
히스토그램, Xbar-R 관리도를 수식과 엑셀 차트로 자동으로 채운다.

사용:  python3 make_spc_template.py [출력파일.xlsx]
"""
import sys

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.chart.layout import Layout, ManualLayout
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.drawing.line import LineProperties
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

SHARE_PATH = r"\\AAS-PTK-CERAMIC\ASKR_Production\01_MFG-line\00. 공정관리\01. Element Line\09. 전반\전반\Chip\선행,양산data\2026"
RESULT_CSV = "공정능력_결과기록.csv"
DATA_CSV = "공정능력_측정데이터.csv"

MAX_LOTS = 30          # SPC 시트에 표시하는 최대 Lot 수 (G ~ AJ 열)
SHOW_ROWS = 20         # 측정 Data 표의 행 수 (Lot 당 측정 수)
HELPER_ROWS = 750      # 계산 시트의 측정값 행 (기록 1건 최대 750개 = 30 Lot x 25)
BINS = 80              # 히스토그램 막대 수

SPC = "SPC Chart"
CALC = "계산"
RES = "결과기록"
DAT = "측정데이터"
GUIDE = "설정 안내"

# 계산 시트에서 결과 CSV 로부터 꺼내오는 항목 (행 번호는 FIELD_ROW 에 기록)
FIELDS = [
    "기록ID", "저장일시", "작성자", "제목(조성)", "Job(Lot)", "고객명", "품명", "공정명", "항목", "작업자", "측정일자",
    "데이터배열", "부분군(개수x크기)", "LSL", "Target", "USL", "N", "평균", "StDev(전체)", "StDev(군내)",
    "Cp", "CPL", "CPU", "Cpk", "Pp", "PPL", "PPU", "Ppk", "Cpm",
    "PPM관측", "PPM기대(전체)", "PPM기대(군내)",
    "Xbar UCL", "Xbar CL", "Xbar LCL", "R UCL", "R CL", "R USL", "Xbar 이탈수", "R 이탈수",
]

thin = Side(style="thin", color="000000")
hair = Side(style="thin", color="B7BCC6")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
GRID = Border(left=hair, right=hair, top=hair, bottom=hair)
HEAD_FILL = PatternFill("solid", fgColor="D9E1F2")
TITLE_FILL = PatternFill("solid", fgColor="D9E1F2")
PICK_FILL = PatternFill("solid", fgColor="FFF2CC")
NOTE_FILL = PatternFill("solid", fgColor="FFFF99")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
RIGHT = Alignment(horizontal="right", vertical="center")
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)


def q(sheet):
    return "'" + sheet + "'"


def add_name(wb, name, ref):
    wb.defined_names[name] = DefinedName(name, attr_text=ref)


def style_range(ws, rng, border=None, fill=None, font=None, align=None):
    for row in ws[rng]:
        for c in row:
            if border: c.border = border
            if fill: c.fill = fill
            if font: c.font = font
            if align: c.alignment = align


def merge(ws, rng, value=None, font=None, fill=None, align=CENTER, border=BOX):
    ws.merge_cells(rng)
    first = ws[rng.split(":")[0]]
    if value is not None:
        first.value = value
    style_range(ws, rng, border=border, fill=fill, font=font, align=align)
    return first


def build(path):
    wb = Workbook()
    spc = wb.active
    spc.title = SPC
    calc = wb.create_sheet(CALC)
    res = wb.create_sheet(RES)
    dat = wb.create_sheet(DAT)
    guide = wb.create_sheet(GUIDE)

    # ------------------------------------------------------------------ 이름 정의
    add_name(wb, "R_HDR", f"{q(RES)}!$1:$1")
    add_name(wb, "R_BODY", f"{q(RES)}!$A$2:$CZ$100000")
    add_name(wb, "R_ID", f"INDEX(R_BODY,0,MATCH(\"기록ID\",R_HDR,0))")
    add_name(wb, "R_LAST", "MATCH(REPT(\"z\",50),R_ID)")
    add_name(wb, "L_ID", "INDEX(R_ID,1):INDEX(R_ID,R_LAST)")
    add_name(wb, "D_HDR", f"{q(DAT)}!$1:$1")
    add_name(wb, "D_BODY", f"{q(DAT)}!$A$2:$J$1000000")
    add_name(wb, "D_ID", "INDEX(D_BODY,0,MATCH(\"기록ID\",D_HDR,0))")
    add_name(wb, "D_LOT", "INDEX(D_BODY,0,MATCH(\"Lot\",D_HDR,0))")
    add_name(wb, "D_SEQ", "INDEX(D_BODY,0,MATCH(\"순번\",D_HDR,0))")
    add_name(wb, "D_VAL", "INDEX(D_BODY,0,MATCH(\"값\",D_HDR,0))")

    # ================================================================== 계산 시트
    calc["A1"] = "선택 기록ID"
    calc["B1"] = f"=IFERROR(IF(TRIM({q(SPC)}!$C$3)<>\"\",TRIM({q(SPC)}!$C$3),INDEX(R_ID,R_LAST)),\"\")"
    calc["A2"] = "결과기록 행"
    calc["B2"] = "=IFERROR(MATCH($B$1,R_ID,0),\"\")"
    calc["A3"] = "측정데이터 첫 행"
    calc["B3"] = "=IFERROR(MATCH($B$1,D_ID,0),0)"
    calc["A4"] = "측정값 수"
    calc["B4"] = "=IF($B$3=0,0,MIN(COUNTIF(D_ID,$B$1)," + str(HELPER_ROWS) + "))"
    calc["A5"] = "Lot 수"
    calc["B5"] = f"=IF($B$4=0,0,MIN(MAX($F$12:$F${11 + HELPER_ROWS}),{MAX_LOTS}))"
    calc["A6"] = "최대값"
    calc["B6"] = f"=IF($B$4=0,\"\",MAX($E$12:$E${11 + HELPER_ROWS}))"
    calc["A7"] = "최소값"
    calc["B7"] = f"=IF($B$4=0,\"\",MIN($E$12:$E${11 + HELPER_ROWS}))"
    for c in "AB":
        calc.column_dimensions[c].width = 16

    # 결과 CSV 항목 (AM: 항목명, AN: 글자 그대로, AO: 숫자면 숫자)
    calc["AM1"] = "항목"
    calc["AN1"] = "값(글자)"
    calc["AO1"] = "값(숫자)"
    field_row = {}
    for i, name in enumerate(FIELDS):
        r = 2 + i
        field_row[name] = r
        calc[f"AM{r}"] = name
        calc[f"AN{r}"] = f"=IF($B$2=\"\",\"\",IFERROR(INDEX(R_BODY,$B$2,MATCH($AM{r},R_HDR,0))&\"\",\"\"))"
        calc[f"AO{r}"] = f"=IF(AN{r}=\"\",\"\",IFERROR(VALUE(AN{r}),AN{r}))"
    calc.column_dimensions["AM"].width = 18
    calc.column_dimensions["AN"].width = 26

    def F(name):            # 계산 시트의 결과 값 (숫자면 숫자)
        return f"{q(CALC)}!$AO${field_row[name]}"

    def T(name):            # 계산 시트의 결과 값 (글자 그대로)
        return f"{q(CALC)}!$AN${field_row[name]}"

    def FL(name):           # 같은 시트 안에서 쓰는 참조
        return f"$AO${field_row[name]}"

    def num(name):          # 숫자가 아니면 #N/A (차트에서 선이 안 그려지도록)
        return f"IF(ISNUMBER({FL(name)}),{FL(name)},NA())"

    # 측정값 (A: i, B: 측정데이터 행, C: Lot, D: 순번, E: 값, F: Lot 번호, G: 키)
    h0 = 11
    for col, title in zip("ABCDEFG", ["i", "행", "Lot", "순번", "값", "Lot번호", "키"]):
        calc[f"{col}{h0}"] = title
    calc[f"A{h0}"].font = Font(bold=True)
    for i in range(1, HELPER_ROWS + 1):
        r = h0 + i
        calc[f"A{r}"] = i
        calc[f"B{r}"] = f"=IF(A{r}<=$B$4,$B$3+A{r}-1,\"\")"
        calc[f"C{r}"] = f"=IF(B{r}=\"\",\"\",INDEX(D_LOT,B{r})&\"\")"
        calc[f"D{r}"] = f"=IF(B{r}=\"\",\"\",INDEX(D_SEQ,B{r}))"
        calc[f"E{r}"] = f"=IF(B{r}=\"\",\"\",IFERROR(VALUE(INDEX(D_VAL,B{r})&\"\"),\"\"))"
        prev = "0" if i == 1 else f"F{r - 1}"
        cond = "1" if i == 1 else f"(C{r}<>C{r - 1})"
        calc[f"F{r}"] = f"=IF(B{r}=\"\",\"\",{prev}+{cond})"
        calc[f"G{r}"] = f"=IF(B{r}=\"\",\"\",F{r}&\"|\"&D{r})"
    HR = f"$12:${11 + HELPER_ROWS}"
    E_RNG = f"$E$12:$E${11 + HELPER_ROWS}"
    F_RNG = f"$F$12:$F${11 + HELPER_ROWS}"
    C_RNG = f"$C$12:$C${11 + HELPER_ROWS}"

    # Lot 표 (J: k, K: Lot, L: 개수, M: 평균, N: 범위, O..T: Xbar 차트, U..Y: R 차트)
    heads = ["k", "Lot", "개수", "평균", "범위", "Xbar", "UCL(+3σ)", "CL", "LCL(-3σ)", "Xbar 이탈", "Lot명",
             "R", "UCL(+3σ)", "CL", "USL", "R 이탈"]
    for j, t in enumerate(heads):
        calc.cell(row=1, column=10 + j, value=t).font = Font(bold=True)
    for k in range(1, MAX_LOTS + 1):
        r = 1 + k
        calc[f"J{r}"] = k
        calc[f"K{r}"] = f"=IFERROR(INDEX({C_RNG},MATCH(J{r},{F_RNG},0)),\"\")"
        calc[f"L{r}"] = f"=COUNTIF({F_RNG},J{r})"
        calc[f"M{r}"] = f"=IF(L{r}=0,\"\",AVERAGEIF({F_RNG},J{r},{E_RNG}))"
        blk = f"INDEX({E_RNG},MATCH(J{r},{F_RNG},0)):INDEX({E_RNG},MATCH(J{r},{F_RNG},0)+L{r}-1)"
        calc[f"N{r}"] = f"=IF(L{r}=0,\"\",MAX({blk})-MIN({blk}))"
        on = f"J{r}<=$B$5"
        calc[f"O{r}"] = f"=IF(AND({on},L{r}>0),M{r},NA())"
        calc[f"P{r}"] = f"=IF({on},{num('Xbar UCL')},NA())"
        calc[f"Q{r}"] = f"=IF({on},{num('Xbar CL')},NA())"
        calc[f"R{r}"] = f"=IF({on},{num('Xbar LCL')},NA())"
        calc[f"S{r}"] = f"=IF(AND({on},L{r}>0),IF(OR(M{r}>P{r},M{r}<R{r}),M{r},NA()),NA())"
        calc[f"T{r}"] = f"=IF({on},K{r},\"\")"
        calc[f"U{r}"] = f"=IF(AND({on},L{r}>0),N{r},NA())"
        calc[f"V{r}"] = f"=IF({on},{num('R UCL')},NA())"
        calc[f"W{r}"] = f"=IF({on},{num('R CL')},NA())"
        calc[f"X{r}"] = f"=IF({on},{num('R USL')},NA())"
        calc[f"Y{r}"] = f"=IF(AND({on},L{r}>0),IF(N{r}>V{r},N{r},NA()),NA())"
    lot_last = 1 + MAX_LOTS

    # 히스토그램 (AB: 막대 시작, AC: 중심, AD: 라벨, AE: 빈도, AF: 전체, AG: 군내, AH: LSL, AI: Target, AJ: USL)
    calc["AA1"] = "히스토그램"
    calc["AA2"] = "시작"
    calc["AA3"] = "막대 너비"
    calc["AA4"] = "세로 최대"
    lo = f"MIN($B$7,IF(ISNUMBER({FL('LSL')}),{FL('LSL')},$B$7))"
    hi = f"MAX($B$6,IF(ISNUMBER({FL('USL')}),{FL('USL')},$B$6))"
    calc["AB2"] = f"=IF($B$4=0,0,{lo}-({hi}-{lo})*0.03)"
    calc["AB3"] = f"=IF($B$4=0,1,MAX(({hi}-{lo})*1.06,1E-9)/{BINS})"
    b0 = 6
    for col, t in zip(["AB", "AC", "AD", "AE", "AF", "AG", "AH", "AI", "AJ"],
                      ["시작", "중심", "라벨", "빈도", "전체", "군내", "LSL", "Target", "USL"]):
        calc[f"{col}{b0}"] = t
    for k in range(1, BINS + 1):
        r = b0 + k
        calc[f"AB{r}"] = f"=$AB$2+({k}-1)*$AB$3"
        calc[f"AC{r}"] = f"=AB{r}+$AB$3/2"
        calc[f"AD{r}"] = f"=TEXT(AC{r},IF($AB$3<0.1,\"0.00\",IF($AB$3<1,\"0.0\",\"0\")))"
        last = " ,\"<=\"" if k == BINS else ",\"<\""
        upper = f"AB{r}+$AB$3"
        calc[f"AE{r}"] = (f"=IF($B$4=0,0,COUNTIFS({q(CALC)}!{E_RNG},\">=\"&AB{r},{q(CALC)}!{E_RNG},"
                          f"\"{'<=' if k == BINS else '<'}\"&{upper}))")
        n_ = f"{FL('N')}"
        calc[f"AF{r}"] = (f"=IF(AND(ISNUMBER({FL('StDev(전체)')}),ISNUMBER({n_})),"
                          f"{n_}*$AB$3*NORMDIST(AC{r},{FL('평균')},{FL('StDev(전체)')},FALSE),NA())")
        calc[f"AG{r}"] = (f"=IF(AND(ISNUMBER({FL('StDev(군내)')}),ISNUMBER({n_})),"
                          f"{n_}*$AB$3*NORMDIST(AC{r},{FL('평균')},{FL('StDev(군내)')},FALSE),NA())")
        for col, name in (("AH", "LSL"), ("AI", "Target"), ("AJ", "USL")):
            calc[f"{col}{r}"] = (f"=IF(ISNUMBER({FL(name)}),IF(AND({FL(name)}>=AB{r},{FL(name)}<AB{r}+$AB$3),$AB$4,NA()),NA())")
    bin_last = b0 + BINS
    calc["AB4"] = f"=MAX(1,MAX($AE${b0 + 1}:$AE${bin_last}),IFERROR(MAX($AF${b0 + 1}:$AG${bin_last}),0))*1.05"

    # ================================================================== SPC Chart 시트
    ws = spc
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 85
    widths = {"A": 6, "B": 9, "C": 13, "D": 1.5, "E": 6, "F": 4}
    for c, w in widths.items():
        ws.column_dimensions[c].width = w
    lot_cols = [get_column_letter(7 + i) for i in range(MAX_LOTS)]   # G .. AJ
    for c in lot_cols:
        ws.column_dimensions[c].width = 8.6
    last_col = lot_cols[-1]

    merge(ws, f"A1:{last_col}2", "SPC Chart", font=Font(size=36, bold=True), fill=TITLE_FILL,
          border=Border(left=Side(style="medium", color="1F3FBF"), right=Side(style="medium", color="1F3FBF"),
                        top=Side(style="medium", color="1F3FBF"), bottom=Side(style="medium", color="1F3FBF")))
    ws.row_dimensions[1].height = 30
    ws.row_dimensions[2].height = 30

    # 기록 선택
    merge(ws, "A3:B3", "기록 선택 ▶", font=Font(bold=True), fill=HEAD_FILL)
    merge(ws, "C3:J3", None, font=Font(bold=True, size=11), fill=PICK_FILL, align=LEFT)
    dv = DataValidation(type="list", formula1="=L_ID", allow_blank=True, showDropDown=False,
                        promptTitle="기록 선택", prompt="목록에서 기록ID 를 고르세요. 비워두면 가장 최근 기록이 표시됩니다.",
                        showErrorMessage=False)
    ws.add_data_validation(dv)
    dv.add("C3")
    merge(ws, "K3:P3", f"=IF({q(CALC)}!$B$1=\"\",\"(표시할 기록 없음 - 설정 안내 시트 참고)\",\"표시 중: \"&{q(CALC)}!$B$1)",
          font=Font(color="1F4E9C", bold=True), align=LEFT, border=None)
    merge(ws, f"Q3:{last_col}3", f"=IF({T('저장일시')}=\"\",\"\",\"저장일시 \"&{T('저장일시')}&\"   작성자 \"&{T('작성자')}&\"   부분군 \"&{T('부분군(개수x크기)')})",
          font=Font(color="555555"), align=LEFT, border=None)

    ws["A4"] = "1. Specification"
    ws["A4"].font = Font(bold=True, size=12)
    ws["E4"] = "3. 측정Data"
    ws["E4"].font = Font(bold=True, size=12)

    # Specification (A~C)
    spec_rows = [
        ("고객명", f"={T('고객명')}"), ("품  명", f"={T('품명')}"), ("조  성", f"={T('제목(조성)')}"),
        ("품번 customer", "-"), ("품번 Amphenol", "-"), ("차  종", ""), ("작업자", f"={T('작업자')}"),
    ]
    r = 5
    for label, val in spec_rows:
        merge(ws, f"A{r}:B{r + 1}", label, fill=HEAD_FILL, font=Font(bold=True))
        merge(ws, f"C{r}:C{r + 1}", val, font=Font(size=11))
        r += 2
    # 규격 (Nominal / +Tol / -Tol / 종류)
    merge(ws, f"A{r}:A{r + 7}", "규\n\n격", fill=HEAD_FILL, font=Font(bold=True))
    nominal = f"IF(ISNUMBER({F('Target')}),{F('Target')},IF(AND(ISNUMBER({F('LSL')}),ISNUMBER({F('USL')})),({F('LSL')}+{F('USL')})/2,\"\"))"
    tol_rows = [
        ("Nominal", f"={nominal}"),
        ("+Tol(USL)", f"=IF(AND(ISNUMBER({F('USL')}),ISNUMBER(C{r})),{F('USL')}-C{r},\"\")"),
        ("-Tol(LSL)", f"=IF(AND(ISNUMBER({F('LSL')}),ISNUMBER(C{r})),C{r}-{F('LSL')},\"\")"),
        ("종  류", f"=IF(AND(ISNUMBER({F('LSL')}),ISNUMBER({F('USL')})),\"양쪽규격\",IF(OR(ISNUMBER({F('LSL')}),ISNUMBER({F('USL')})),\"한쪽규격\",\"\"))"),
    ]
    for label, val in tol_rows:
        merge(ws, f"B{r}:B{r + 1}", label, fill=HEAD_FILL, font=Font(size=9))
        merge(ws, f"C{r}:C{r + 1}", val, font=Font(size=11))
        r += 2
    spec_end = r - 1   # 26

    # 측정 Data 머리 (5행: 공정명 / 항목 / 측정일자, 6행: Lot)
    merge(ws, "E5:F5", "공정명", fill=HEAD_FILL, font=Font(bold=True))
    merge(ws, "G5:S5", f"={T('공정명')}", font=Font(size=16, bold=True))
    merge(ws, "T5:V5", "항  목", fill=HEAD_FILL, font=Font(bold=True))
    merge(ws, "W5:AB5", f"={T('항목')}", font=Font(size=12))
    merge(ws, "AC5:AE5", "측정 일자", fill=HEAD_FILL, font=Font(bold=True))
    merge(ws, f"AF5:{last_col}5", f"={T('측정일자')}", font=Font(size=12))
    ws.row_dimensions[5].height = 24
    merge(ws, "E6:F6", "Lot Date", fill=HEAD_FILL, font=Font(bold=True))
    for k, col in enumerate(lot_cols, start=1):
        c = ws[f"{col}6"]
        c.value = f"={q(CALC)}!$T${1 + k}&\"\""
        c.font = Font(size=9, bold=True)
        c.alignment = CENTER
        c.border = BOX
        c.fill = HEAD_FILL
    # 측정값 7..26
    for j in range(1, SHOW_ROWS + 1):
        rr = 6 + j
        merge(ws, f"E{rr}:F{rr}", j, font=Font(size=10))
        for k, col in enumerate(lot_cols, start=1):
            c = ws[f"{col}{rr}"]
            c.value = (f"=IF({col}$6=\"\",\"\",IFERROR(INDEX({q(CALC)}!{E_RNG},"
                       f"MATCH(\"{k}|{j}\",{q(CALC)}!$G$12:$G${11 + HELPER_ROWS},0)),\"\"))")
            c.number_format = "0.0"
            c.alignment = CENTER
            c.border = GRID
            c.font = Font(size=10)
    r_avg, r_rng, r_ok = 7 + SHOW_ROWS, 8 + SHOW_ROWS, 9 + SHOW_ROWS   # 27, 28, 29
    merge(ws, f"E{r_avg}:F{r_avg}", "평균(X)", fill=HEAD_FILL, font=Font(bold=True, size=9))
    merge(ws, f"E{r_rng}:F{r_rng}", "범위(R)", fill=HEAD_FILL, font=Font(bold=True, size=9))
    merge(ws, f"E{r_ok}:F{r_ok}", "관리 상태", fill=HEAD_FILL, font=Font(bold=True, size=9))
    for k, col in enumerate(lot_cols, start=1):
        cr = 1 + k
        a = ws[f"{col}{r_avg}"]
        a.value = f"=IF({col}$6=\"\",\"\",{q(CALC)}!$M${cr})"
        a.number_format = "0.00"
        b = ws[f"{col}{r_rng}"]
        b.value = f"=IF({col}$6=\"\",\"\",{q(CALC)}!$N${cr})"
        b.number_format = "0.000"
        o = ws[f"{col}{r_ok}"]
        o.value = (f"=IF({col}$6=\"\",\"\",IF(OR(ISNUMBER({q(CALC)}!$S${cr}),ISNUMBER({q(CALC)}!$Y${cr})),\"이탈\",\"OK\"))")
        for cell in (a, b, o):
            cell.alignment = CENTER
            cell.border = BOX
            cell.font = Font(size=10, bold=True)
    # 이탈 강조 (빨간 글씨)
    red = Font(color="D00000", bold=True)
    rng_lots = f"G{r_avg}:{last_col}{r_avg}"
    ws.conditional_formatting.add(rng_lots, FormulaRule(formula=[f"ISNUMBER(INDEX({q(CALC)}!$S$2:$S${lot_last},COLUMN()-6))"], font=red))
    ws.conditional_formatting.add(f"G{r_rng}:{last_col}{r_rng}", FormulaRule(formula=[f"ISNUMBER(INDEX({q(CALC)}!$Y$2:$Y${lot_last},COLUMN()-6))"], font=red))
    ws.conditional_formatting.add(f"G{r_ok}:{last_col}{r_ok}", FormulaRule(formula=[f"G{r_ok}=\"이탈\""], font=red,
                                  fill=PatternFill("solid", fgColor="FFE1E1")))

    # ------------------------------------------------------------------ 2. 통계치 Summary
    s0 = r_ok + 2   # 31
    ws[f"A{s0}"] = "2. 통계치 Summary"
    ws[f"A{s0}"].font = Font(bold=True, size=12)
    ws[f"E{s0}"] = "4. Histogram"
    ws[f"E{s0}"].font = Font(bold=True, size=12)
    ws[f"S{s0}"] = "5. X̄-R 관리도"
    ws[f"S{s0}"].font = Font(bold=True, size=12)
    cpk = F("Cpk")
    summary = [
        ("추정불량율 /ppm", f"={F('PPM기대(군내)')}", "0.00", False),
        ("Cpk", f"={cpk}", "0.00", True),
        ("Cpk 판정", f"=IF(ISNUMBER({cpk}),IF({cpk}>=1.33,\"OK (1.33이상)\",\"NG (1.33미만)\"),\"\")", None, True),
        ("Cpu", f"={F('CPU')}", "0.00", False),
        ("Cpl", f"={F('CPL')}", "0.00", False),
        ("Cp", f"={F('Cp')}", "0.00", True),
        ("Ppk", f"={F('Ppk')}", "0.00", False),
        ("최대값", f"={q(CALC)}!$B$6", "0.000", False),
        ("최소값", f"={q(CALC)}!$B$7", "0.000", False),
        ("평균(X)", f"={F('평균')}", "0.000", False),
        ("표준편차(전체)", f"={F('StDev(전체)')}", "0.0000", False),
        ("표준편차(군내)", f"={F('StDev(군내)')}", "0.0000", False),
    ]
    r = s0 + 1
    for label, val, fmt, bold in summary:
        merge(ws, f"A{r}:B{r}", label, fill=HEAD_FILL if bold else None, font=Font(bold=bold, size=10))
        c = ws[f"C{r}"]
        c.value = val
        c.border = BOX
        c.alignment = RIGHT
        c.font = Font(bold=bold, size=11)
        if fmt:
            c.number_format = fmt
        r += 1
    limits = [("X", "UCL", "Xbar UCL"), ("X", "CL", "Xbar CL"), ("X", "LCL", "Xbar LCL"),
              ("R", "UCL", "R UCL"), ("R", "CL", "R CL"), ("R", "USL", "R USL")]
    x_start = r
    for grp, label, field in limits:
        ws[f"B{r}"] = label
        ws[f"B{r}"].border = BOX
        ws[f"B{r}"].alignment = CENTER
        c = ws[f"C{r}"]
        c.value = f"={F(field)}"
        c.number_format = "#,##0.00"
        c.border = BOX
        c.alignment = RIGHT
        r += 1
    merge(ws, f"A{x_start}:A{x_start + 2}", "X", fill=HEAD_FILL, font=Font(bold=True))
    merge(ws, f"A{x_start + 3}:A{x_start + 5}", "R", fill=HEAD_FILL, font=Font(bold=True))
    merge(ws, f"A{r}:B{r}", "이탈 Lot (X / R)", font=Font(size=10))
    ws[f"C{r}"] = f"=IF({T('Xbar 이탈수')}=\"\",\"\",{T('Xbar 이탈수')}&\" / \"&{T('R 이탈수')})"
    ws[f"C{r}"].border = BOX
    ws[f"C{r}"].alignment = RIGHT
    summary_end = r

    # 판정 기준
    crit_row = max(summary_end, s0 + 24) + 2
    merge(ws, f"A{crit_row}:F{crit_row}", "★ 공정능력 유무의 판단 기준 ★", fill=NOTE_FILL, font=Font(bold=True), border=None)
    merge(ws, f"G{crit_row}:{last_col}{crit_row}",
          "▶ Cpk ≥ 1.67 : 공정능력 매우 충분     ▶ 1.67 > Cpk ≥ 1.33 : 공정능력이 충분     ▶ 1.33 > Cpk ≥ 1.00 : 공정능력이 있음     "
          "▶ 1.00 > Cpk ≥ 0.67 : 공정능력이 부족     ▶ 0.67 > Cpk : 공정능력이 없음",
          fill=NOTE_FILL, font=Font(size=10), align=LEFT, border=None)

    # ------------------------------------------------------------------ 차트
    def series_style(s, color, width=19050, dash=None, marker=None, marker_color=None, no_line=False, size=6):
        s.smooth = False
        gp = s.graphicalProperties
        if no_line:
            gp.line.noFill = True
        else:
            gp.line.solidFill = color
            gp.line.width = width
            if dash:
                gp.line.dashStyle = dash
        if marker:
            s.marker.symbol = marker
            s.marker.size = size
            s.marker.graphicalProperties = GraphicalProperties(solidFill=marker_color or color)
            s.marker.graphicalProperties.line.solidFill = marker_color or color
        else:
            s.marker.symbol = "none"

    def finish_axes(ch, y_title=None):
        ch.display_blanks = "gap"      # #N/A / 빈 칸은 점을 찍지 않음
        ch.x_axis.delete = False
        ch.y_axis.delete = False
        ch.legend.position = "b"
        if y_title:
            ch.y_axis.title = y_title

    calc_ref = lambda c1, r1, c2, r2: Reference(calc, min_col=c1, min_row=r1, max_col=c2, max_row=r2)
    col = lambda letters: __import__("openpyxl").utils.column_index_from_string(letters)

    # Xbar Chart
    xc = LineChart()
    xc.title = "Xbar Chart"
    xc.style = 2
    xc.height = 8.2
    xc.width = 30
    specs = [("O", "Xbar", "1F4E9C", None, "circle", None, False),
             ("P", "UCL(+3σ)", "C00000", "dash", None, None, False),
             ("Q", "CL", "008A3E", None, None, None, False),
             ("R", "LCL(-3σ)", "C00000", "dash", None, None, False),
             ("S", "이탈", "D00000", None, "square", "D00000", True)]
    for letter, _, color, dash, marker, mcol, no_line in specs:
        xc.add_data(calc_ref(col(letter), 1, col(letter), lot_last), titles_from_data=True)
    for s, (letter, _, color, dash, marker, mcol, no_line) in zip(xc.series, specs):
        series_style(s, color, dash=dash, marker=marker, marker_color=mcol, no_line=no_line, size=7 if marker == "square" else 5)
    xc.set_categories(calc_ref(col("T"), 2, col("T"), lot_last))
    finish_axes(xc)
    xc.y_axis.number_format = "0.0"
    xc.x_axis.tickLblSkip = 1
    xc.x_axis.txPr = None
    ws.add_chart(xc, f"S{s0 + 1}")

    # R Chart
    rc = LineChart()
    rc.title = "R Chart"
    rc.style = 2
    rc.height = 8.2
    rc.width = 30
    rspecs = [("U", "R", "1F4E9C", None, "circle", None, False),
              ("V", "UCL(+3σ)", "C00000", "dash", None, None, False),
              ("W", "CL", "008A3E", None, None, None, False),
              ("X", "USL", "E67E22", "sysDot", None, None, False),
              ("Y", "이탈", "D00000", None, "square", "D00000", True)]
    for letter, *_ in rspecs:
        rc.add_data(calc_ref(col(letter), 1, col(letter), lot_last), titles_from_data=True)
    for s, (letter, _, color, dash, marker, mcol, no_line) in zip(rc.series, rspecs):
        series_style(s, color, dash=dash, marker=marker, marker_color=mcol, no_line=no_line, size=7 if marker == "square" else 5)
    rc.set_categories(calc_ref(col("T"), 2, col("T"), lot_last))
    finish_axes(rc)
    rc.y_axis.number_format = "0.0"
    ws.add_chart(rc, f"S{s0 + 18}")

    # Histogram (막대 + 정규분포 곡선 + 규격선)
    hb = BarChart()
    hb.type = "col"
    hb.grouping = "clustered"
    hb.overlap = 100
    hb.gapWidth = 0
    hb.title = "Histogram"
    hb.style = 2
    hb.height = 16.6
    hb.width = 18.5
    hb.add_data(calc_ref(col("AE"), b0, col("AE"), bin_last), titles_from_data=True)
    for letter in ("AH", "AI", "AJ"):
        hb.add_data(calc_ref(col(letter), b0, col(letter), bin_last), titles_from_data=True)
    bar_colors = ["BCCDE5", "C00000", "00A651", "C00000"]
    for s, color in zip(hb.series, bar_colors):
        s.graphicalProperties = GraphicalProperties(solidFill=color)
        s.graphicalProperties.line.solidFill = "000000" if color == "BCCDE5" else color
        s.graphicalProperties.line.width = 6350
    hb.set_categories(calc_ref(col("AD"), b0 + 1, col("AD"), bin_last))
    hl = LineChart()
    hl.add_data(calc_ref(col("AF"), b0, col("AF"), bin_last), titles_from_data=True)
    hl.add_data(calc_ref(col("AG"), b0, col("AG"), bin_last), titles_from_data=True)
    series_style(hl.series[0], "B00000", width=25400)
    series_style(hl.series[1], "404040", width=25400, dash="dash")
    hb.x_axis.delete = False
    hb.y_axis.delete = False
    hb.x_axis.tickLblSkip = 10
    hb.x_axis.tickMarkSkip = 10
    hb.y_axis.title = "빈도수 (Frequency)"
    hb.x_axis.title = "측정값 (Value)"
    hb.legend.position = "t"
    hb.display_blanks = "gap"
    hb += hl
    ws.add_chart(hb, f"E{s0 + 1}")

    ws.freeze_panes = "A4"
    ws.print_area = f"A1:{last_col}{crit_row}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True

    # ================================================================== 설정 안내
    g = guide
    g.column_dimensions["A"].width = 4
    g.column_dimensions["B"].width = 120
    lines = [
        ("SPC Chart 설정 안내 (처음 한 번만)", Font(bold=True, size=16)),
        ("", None),
        ("이 파일은 공용 폴더의 두 CSV 를 '결과기록' / '측정데이터' 시트로 불러와 'SPC Chart' 시트를 자동으로 채웁니다.", None),
        ("CSV 는 공정능력 HTML 프로그램의 [결과 저장] 버튼으로 쌓입니다.  폴더: " + SHARE_PATH, None),
        ("", None),
        ("1단계. '결과기록' 연결", Font(bold=True, size=12)),
        ("  ① 메뉴 [데이터] → [새 쿼리] → [기타 원본에서] → [빈 쿼리]", None),
        ("  ② 열린 파워 쿼리 편집기에서 [보기] → [고급 편집기] → 내용을 모두 지우고 아래 '결과기록 쿼리' 코드를 붙여넣기 → [완료]", None),
        ("  ③ 오른쪽 '쿼리 설정'의 이름을 결과기록 으로 바꾸기", None),
        ("  ④ [홈] → [닫기 및 다음으로 로드] → '표' 선택, '기존 워크시트' 선택 → 결과기록 시트의 A1 셀 클릭 → [확인]", None),
        ("", None),
        ("2단계. '측정데이터' 연결", Font(bold=True, size=12)),
        ("  1단계와 같은 방법으로 아래 '측정데이터 쿼리' 코드를 붙여넣고, 이름은 측정데이터, 로드 위치는 측정데이터 시트의 A1 셀", None),
        ("", None),
        ("3단계. 자동 새로 고침 (두 쿼리 모두)", Font(bold=True, size=12)),
        ("  [데이터] → [연결] (또는 [쿼리 및 연결]) → 쿼리 오른쪽 클릭 → [속성]", None),
        ("  → '다음 간격으로 새로 고침' 체크, 1 분 /  '파일을 열 때 데이터 새로 고침' 체크 → [확인]", None),
        ("", None),
        ("사용 방법", Font(bold=True, size=12)),
        ("  · 'SPC Chart' 시트 C3 칸(노란색)에서 기록ID 를 고르면 그 기록이 표시됩니다. 비워두면 가장 최근 기록이 표시됩니다.", None),
        ("  · 새 기록이 저장되면 1분 안에 자동으로 반영됩니다. 바로 보려면 [데이터] → [모두 새로 고침].", None),
        ("  · 한 기록에 Lot 은 최대 30개, 측정 Data 표에는 Lot 당 20개까지 표시됩니다 (통계·관리도 값은 HTML 이 계산한 값을 그대로 씀).", None),
        ("  · 결과기록 / 측정데이터 / 계산 시트는 직접 고치지 마세요. (새로 고침 때 덮어써집니다)", None),
        ("  · CSV 파일을 엑셀로 직접 열어두면 HTML 이 저장하지 못합니다. 이 SPC 파일은 열어두어도 괜찮습니다.", None),
        ("", None),
        ("결과기록 쿼리", Font(bold=True, size=12, color="1F4E9C")),
    ]
    m_result = (
        "let\n"
        f"    경로 = \"{SHARE_PATH}\\{RESULT_CSV}\",\n"
        "    원본 = Csv.Document(File.Contents(경로), [Delimiter=\",\", Encoding=65001, QuoteStyle=QuoteStyle.Csv]),\n"
        "    머리글 = Table.PromoteHeaders(원본, [PromoteAllScalars=true])\n"
        "in\n"
        "    머리글"
    )
    m_data = (
        "let\n"
        f"    경로 = \"{SHARE_PATH}\\{DATA_CSV}\",\n"
        "    원본 = Csv.Document(File.Contents(경로), [Delimiter=\",\", Encoding=65001, QuoteStyle=QuoteStyle.Csv]),\n"
        "    머리글 = Table.PromoteHeaders(원본, [PromoteAllScalars=true]),\n"
        "    형식 = Table.TransformColumnTypes(머리글, {{\"순번\", Int64.Type}, {\"값\", type number}}, \"en-US\")\n"
        "in\n"
        "    형식"
    )
    r = 1
    for text, font in lines:
        c = g[f"B{r}"]
        c.value = text
        if font:
            c.font = font
        c.alignment = Alignment(wrap_text=True, vertical="top")
        r += 1
    code_font = Font(name="Consolas", size=10)
    code_fill = PatternFill("solid", fgColor="F3F4F6")
    c = g[f"B{r}"]
    c.value = m_result
    c.font = code_font
    c.fill = code_fill
    c.alignment = Alignment(wrap_text=True, vertical="top")
    g.row_dimensions[r].height = 95
    r += 2
    g[f"B{r}"] = "측정데이터 쿼리"
    g[f"B{r}"].font = Font(bold=True, size=12, color="1F4E9C")
    r += 1
    c = g[f"B{r}"]
    c.value = m_data
    c.font = code_font
    c.fill = code_fill
    c.alignment = Alignment(wrap_text=True, vertical="top")
    g.row_dimensions[r].height = 110

    # 시트 정리
    for sh, color in ((res, "A9D08E"), (dat, "A9D08E"), (calc, "BFBFBF"), (guide, "FFC000")):
        sh.sheet_properties.tabColor = color
    spc.sheet_properties.tabColor = "1F3FBF"
    wb.calculation.fullCalcOnLoad = True
    wb.active = 0
    wb.save(path)
    return field_row


if __name__ == "__main__":
    out = sys.argv[1] if len(sys.argv) > 1 else "SPC_Chart.xlsx"
    build(out)
    print("saved", out)
