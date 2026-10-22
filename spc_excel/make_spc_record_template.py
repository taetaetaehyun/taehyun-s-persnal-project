# -*- coding: utf-8 -*-
"""
기록별 SPC 엑셀 템플릿 생성기

기존 SPC 파일(127K_S-SC04C-_B__SPC_2026.xlsx)의 'SC04C- 1033Y~1060Y' 시트를 그대로 복사해
(행 높이, 열 너비, 글꼴, 테두리, 병합, 인쇄 설정, 제목 도형) 값이 들어갈 칸에 {{키}} 자리표시를 넣고,
그래프 자리에 사진 2장(히스토그램 10x13cm, Xbar-R 24x13cm)을 넣은 템플릿을 만든다.

공정능력 HTML 은 [결과 저장] 때 이 템플릿의 자리표시를 값으로, 사진을 실제 그래프 PNG 로 바꿔
공용 폴더 SPC 폴더에 기록별 SPC 파일을 만든다.

사용:  python3 make_spc_record_template.py 원본SPC.xlsx 출력템플릿.xlsx [HTML파일]
       (HTML 파일을 주면 템플릿을 base64 로 HTML 안에 넣어줌)
"""
import base64
import io
import re
import sys
import zipfile

import openpyxl
from openpyxl.drawing.image import Image as XLImage
from PIL import Image

SRC_SHEET = "SC04C- 1033Y~1060Y"
LOT_COLS = 25            # G ~ AE
FIRST_LOT_COL = 7        # G
DATA_ROWS = 20           # 7 ~ 26
EMU_PER_CM = 360000
HIST_SIZE = (10, 13)     # 너비, 높이 (cm)
XBR_SIZE = (24, 13)


def col(n):
    return openpyxl.utils.get_column_letter(n)


def placeholder_png(w_cm, h_cm, text):
    im = Image.new("RGB", (int(w_cm * 40), int(h_cm * 40)), "white")
    buf = io.BytesIO()
    im.save(buf, "PNG")
    buf.seek(0)
    return buf


def build(src, out):
    wb = openpyxl.load_workbook(src)
    ws = wb[SRC_SHEET]
    for other in [s for s in wb.worksheets if s is not ws]:
        wb.remove(other)
    ws.title = "SPC Chart"
    for name in list(wb.defined_names.keys()):
        del wb.defined_names[name]
    ws.defined_names.clear()          # 원본에 딸려 있던 이름 정의(외부 연결 등) 제거
    wb._external_links = []
    ws._charts = []                   # 원본 차트 제거 (사진으로 대체)
    ws._images = []
    ws.print_area = "A1:AE48"

    P = lambda key: "{{" + key + "}}"

    # 1. Specification
    for addr, key in [("C5", "CUSTOMER"), ("C6", "PRODUCT"), ("C7", "PN_CUSTOMER"), ("C10", "PN_AMPHENOL"),
                      ("C13", "CAR"), ("C16", "WORKER"), ("C21", "NOMINAL"), ("C23", "TOL_U"), ("C25", "TOL_L"),
                      ("C27", "SPEC_TYPE"), ("H5", "PROCESS"), ("U5", "ITEM"), ("AB5", "MDATE")]:
        ws[addr] = P(key)

    # 3. 측정 Data (Lot 25 열 x 20 행), 평균 / 범위 / 관리자 확인
    for k in range(1, LOT_COLS + 1):
        c = col(FIRST_LOT_COL + k - 1)
        ws[f"{c}6"] = P(f"LOT_{k}")
        for j in range(1, DATA_ROWS + 1):
            ws[f"{c}{6 + j}"] = P(f"V_{k}_{j}")
            ws[f"{c}{6 + j}"].number_format = "0.0"
        ws[f"{c}27"] = P(f"AVG_{k}")
        ws[f"{c}27"].number_format = "0"
        ws[f"{c}28"] = P(f"RNG_{k}")
        ws[f"{c}29"] = P(f"OK_{k}")

    # 2. 통계치 Summary
    for addr, key in [("C31", "PPM"), ("C32", "CPK"), ("C34", "CPU"), ("C35", "CPL"), ("C36", "CP"),
                      ("C37", "MAX"), ("C38", "MIN"), ("C39", "MEAN"), ("C40", "SIGMA"),
                      ("C41", "X_UCL"), ("C42", "X_CL"), ("C43", "X_LCL"), ("C44", "R_UCL"), ("C45", "R_CL")]:
        ws[addr] = P(key)
    for addr in ("C41", "C42", "C43"):
        ws[addr].number_format = "#,##0.00"
    for addr in ("C44", "C45"):
        ws[addr].number_format = "0.00"
    ws["C40"].number_format = "0.000"

    # 예전 차트용 계산 칸(그래프 뒤에 숨어 있던 값) 비우기
    for r in range(31, 46):
        for cidx in range(5, 36):
            ws.cell(row=r, column=cidx).value = None

    # 그래프 사진 2장 (히스토그램: E31, Xbar-R: M31)
    for anchor, (w, h), name in (("E31", HIST_SIZE, "hist"), ("M31", XBR_SIZE, "xbarr")):
        img = XLImage(placeholder_png(w, h, name))
        img.width = w * EMU_PER_CM / 9525
        img.height = h * EMU_PER_CM / 9525
        img.anchor = anchor
        ws.add_image(img)

    tmp = io.BytesIO()
    wb.save(tmp)

    # 원본 제목 도형(SPC Chart) 을 새 drawing 에 옮겨 넣고, 사진 크기를 cm 정확히 맞춤
    src_zip = zipfile.ZipFile(src)
    src_draw = src_zip.read("xl/drawings/drawing2.xml").decode("utf-8")
    title = None
    for m in re.finditer(r"<xdr:twoCellAnchor[^>]*>.*?</xdr:twoCellAnchor>", src_draw, re.S):
        if "SPC Chart" in m.group(0):
            title = m.group(0)
    assert title, "제목 도형을 찾지 못함"

    zin = zipfile.ZipFile(io.BytesIO(tmp.getvalue()))
    zout_buf = io.BytesIO()
    zout = zipfile.ZipFile(zout_buf, "w", zipfile.ZIP_DEFLATED)
    for item in zin.infolist():
        data = zin.read(item.filename)
        if item.filename.startswith("xl/drawings/drawing") and item.filename.endswith(".xml"):
            x = data.decode("utf-8")
            sizes = iter([HIST_SIZE, XBR_SIZE])

            def fix_ext(m):
                w, h = next(sizes)
                cx, cy = int(w * EMU_PER_CM), int(h * EMU_PER_CM)
                block = re.sub(r'<ext cx="\d+" cy="\d+"\s*/>', f'<ext cx="{cx}" cy="{cy}"/>', m.group(0))
                return re.sub(r'<a:ext cx="\d+" cy="\d+"\s*/>', f'<a:ext cx="{cx}" cy="{cy}"/>', block)

            x = re.sub(r'<oneCellAnchor>.*?</oneCellAnchor>', fix_ext, x, flags=re.S)
            ns = {"xdr": "http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing",
                  "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
                  "a14": "http://schemas.microsoft.com/office/drawing/2010/main"}
            decl = "".join(f' xmlns:{k}="{v}"' for k, v in ns.items() if f"xmlns:{k}=" not in title)
            x = x.replace("</wsDr>", title.replace("<xdr:twoCellAnchor", "<xdr:twoCellAnchor" + decl, 1) + "</wsDr>")
            data = x.encode("utf-8")
        elif item.filename == "xl/worksheets/sheet1.xml":
            data = fix_data_cf(data.decode("utf-8")).encode("utf-8")
        zout.writestr(item, data)
    zout.close()
    open(out, "wb").write(zout_buf.getvalue())
    return zout_buf.getvalue()


def fix_data_cf(x):
    """측정 Data 칸(G7:AE26)의 조건부 서식을 '규격(LSL~USL) 밖이면 빨간 굵은 글씨' 한 가지로 통일.
    원본에는 열마다 고정값(3495~3637), MAX/MIN, 빈칸 참조 등 서로 다른 규칙이 남아 있었음."""
    blocks = re.findall(r"<conditionalFormatting .*?</conditionalFormatting>", x, re.S)
    data_blocks = [b for b in blocks if re.search(r'sqref="[^"]*\b[A-Z]{1,2}7:', b)]
    red = re.search(r'dxfId="(\d+)"', [b for b in blocks if 'sqref="C32"' in b][0]).group(1)
    for b in data_blocks:
        x = x.replace(b, "", 1)
    v = "G7"
    f = (f"AND(ISNUMBER({v}),ISNUMBER($C$21),OR(AND(ISNUMBER($C$25),{v}&lt;$C$21-$C$25),"
         f"AND(ISNUMBER($C$23),{v}&gt;$C$21+$C$23)))")
    rule = (f'<conditionalFormatting sqref="G7:AE26"><cfRule type="expression" priority="1" dxfId="{red}">'
            f"<formula>{f}</formula></cfRule></conditionalFormatting>")
    i = x.index("<conditionalFormatting ")
    return x[:i] + rule + x[i:]


def embed(html_path, data):
    html = open(html_path, encoding="utf-8").read()
    b64 = base64.b64encode(data).decode("ascii")
    new = f"const SPC_TEMPLATE_B64 = '{b64}';"
    html2, n = re.subn(r"const SPC_TEMPLATE_B64 = '[^']*';", new, html)
    assert n == 1, "HTML 안에 SPC_TEMPLATE_B64 자리가 없음"
    open(html_path, "w", encoding="utf-8").write(html2)
    print("embedded", len(b64), "chars into", html_path)


if __name__ == "__main__":
    data = build(sys.argv[1], sys.argv[2])
    print("saved", sys.argv[2], len(data), "bytes")
    if len(sys.argv) > 3:
        embed(sys.argv[3], data)
